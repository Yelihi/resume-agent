import asyncio
from io import BytesIO

import httpx
import pymupdf
from docx import Document

import app.resume.api as resume_api
from app.document_processing.models import (
    DocumentPage,
    NormalizedBBox,
    PageBlock,
    PageDocument,
    PageLine,
)
from app.main import app


def request(method: str, path: str, **kwargs) -> httpx.Response:
    async def send() -> httpx.Response:
        transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            return await client.request(method, path, **kwargs)

    return asyncio.run(send())


def test_returns_validation_policy() -> None:
    response = request("GET", "/api/resumes/validation-policy")

    assert response.status_code == 200
    assert set(response.json()["maximumFileSizeBytes"]) == {"pdf", "image", "docx", "txt"}
    assert response.json()["maximumDirectTextCharacters"] > 0


def test_extracts_direct_text_and_txt() -> None:
    direct = request("POST", "/api/resumes/extract/text", json={"text": "A😀B\nsecond"})
    uploaded = request(
        "POST",
        "/api/resumes/extract/txt",
        files={"file": ("resume.txt", "one\ntwo".encode(), "text/plain")},
    )

    assert direct.status_code == uploaded.status_code == 200
    assert direct.json()["document"]["blocks"][0]["lines"][0]["endOffset"] == 4
    assert uploaded.json()["document"]["text"] == "one\ntwo"


def test_extracts_docx() -> None:
    source = Document()
    source.add_paragraph("Experience")
    output = BytesIO()
    source.save(output)

    response = request(
        "POST",
        "/api/resumes/extract/docx",
        files={"file": ("resume.docx", output.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 200
    assert response.json()["document"]["text"] == "Experience"


def test_extracts_text_pdf() -> None:
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text((72, 72), "Experience")

    response = request(
        "POST",
        "/api/resumes/extract/pdf",
        files={"file": ("resume.pdf", source.tobytes(), "application/pdf")},
    )

    assert response.status_code == 200
    assert response.json()["document"]["pages"][0]["blocks"][0]["lines"][0]["text"] == "Experience"


def test_extracts_image_with_fixed_page_response(monkeypatch) -> None:
    box = NormalizedBBox(x=0.1, y=0.1, width=0.5, height=0.1)
    document = PageDocument(
        pages=[
            DocumentPage(
                pageNumber=1,
                blocks=[
                    PageBlock(
                        blockId="p1-b1",
                        bbox=box,
                        lines=[
                            PageLine(
                                lineId="p1-l1",
                                text="Image resume",
                                bbox=box,
                                textSource="ocr",
                                uncertainWords=[],
                            )
                        ],
                    )
                ],
            )
        ]
    )
    monkeypatch.setattr(resume_api, "extract_image", lambda _: document)

    response = request(
        "POST",
        "/api/resumes/extract/image",
        files={"file": ("resume.png", b"\x89PNG\r\n\x1a\n", "image/png")},
    )

    assert response.status_code == 200
    assert response.json()["document"]["pages"][0]["blocks"][0]["lines"][0]["text"] == "Image resume"


def test_rejects_disguised_format_with_safe_error_contract() -> None:
    response = request(
        "POST",
        "/api/resumes/extract/pdf",
        files={"file": ("resume.pdf", b"plain text", "application/pdf")},
    )

    assert response.status_code == 415
    assert response.json()["errors"][0] == {
        "moduleKey": "resumeInputValidation",
        "inputSourceId": None,
        "errorCode": "UNSUPPORTED_MEDIA_TYPE",
        "userMessage": "선택한 입력 형식과 실제 파일 형식이 다릅니다.",
        "canRetry": False,
    }


def test_rejects_file_over_policy_limit(monkeypatch) -> None:
    monkeypatch.setattr(resume_api.POLICY.maximumFileSizeBytes, "txt", 2)

    response = request(
        "POST",
        "/api/resumes/extract/txt",
        files={"file": ("resume.txt", b"long", "text/plain")},
    )

    assert response.status_code == 413
    assert response.json()["errors"][0]["errorCode"] == "FILE_TOO_LARGE"


def test_maps_corrupt_document_and_unexpected_exception(monkeypatch) -> None:
    corrupt = request(
        "POST",
        "/api/resumes/extract/pdf",
        files={"file": ("resume.pdf", b"%PDF-1.7\nbroken", "application/pdf")},
    )
    assert corrupt.status_code == 422
    assert corrupt.json()["errors"][0]["errorCode"] == "CORRUPT_PDF"

    def fail(_: bytes):
        raise RuntimeError("sensitive internal detail")

    monkeypatch.setattr(resume_api, "extract_txt", fail)
    unexpected = request(
        "POST",
        "/api/resumes/extract/txt",
        files={"file": ("resume.txt", b"valid", "text/plain")},
    )
    assert unexpected.status_code == 500
    assert unexpected.json()["errors"][0]["errorCode"] == "INTERNAL_ERROR"
    assert "sensitive" not in unexpected.text

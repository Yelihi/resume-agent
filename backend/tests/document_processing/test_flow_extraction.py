from io import BytesIO

from docx import Document

from app.document_processing.flow import extract_docx, extract_text, extract_txt


def test_text_normalizes_only_bom_and_newlines_and_splits_blocks() -> None:
    document = extract_text("\ufeff첫 줄\r\nA😀B\r\n  \r마지막  ")

    assert document.text == "첫 줄\nA😀B\n  \n마지막  "
    assert [block.blockId for block in document.blocks] == ["f-b1", "f-b2"]
    assert [line.text for block in document.blocks for line in block.lines] == [
        "첫 줄",
        "A😀B",
        "마지막  ",
    ]
    assert [line.lineId for block in document.blocks for line in block.lines] == [
        "f-l1",
        "f-l2",
        "f-l3",
    ]
    assert document.blocks[0].lines[1].endOffset - document.blocks[0].lines[1].startOffset == 4


def test_txt_uses_same_flow_extraction() -> None:
    raw = "\ufeffone\r\ntwo".encode("utf-8")

    assert extract_txt(raw) == extract_text("one\ntwo")


def test_docx_preserves_header_body_table_and_footer_order() -> None:
    source = Document()
    source.sections[0].header.paragraphs[0].text = "HEADER"
    source.add_paragraph("Intro")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "React"
    table.cell(0, 1).text = "TypeScript"
    source.add_paragraph("End")
    source.sections[0].footer.paragraphs[0].text = "FOOTER"
    output = BytesIO()
    source.save(output)

    document = extract_docx(output.getvalue())

    assert document.text == "HEADER\n\nIntro\nReact\tTypeScript\nEnd\n\nFOOTER"
    assert [line.text for block in document.blocks for line in block.lines] == [
        "HEADER",
        "Intro",
        "React\tTypeScript",
        "End",
        "FOOTER",
    ]


def test_every_flow_line_resolves_from_utf16_offsets() -> None:
    document = extract_text("가😀나\nsecond")
    encoded = document.text.encode("utf-16-le")

    for block in document.blocks:
        for line in block.lines:
            resolved = encoded[line.startOffset * 2 : line.endOffset * 2].decode("utf-16-le")
            assert resolved == line.text
